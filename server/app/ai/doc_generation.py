"""Document generation capability for TianShu.

This module backs the ``doc-tools`` skill registered in
``TianShuSkillRegistry``. It turns the *live* runtime scenario snapshot into a
real situation report and writes it to the shared document output directory
(``TIANSHU_DOC_OUTPUT_DIR``, default ``/doc_output``) so the pre-existing
``/api/doc/*`` routes can serve / preview / list the result.

The capability is intentionally framework-neutral: it only depends on the
``TianShuRuntime`` (to read the current scenario) and ``python-docx`` (optional
-- a Markdown companion file is always produced so a readable artifact exists
even when ``python-docx`` is unavailable).
"""

from __future__ import annotations

import os
import re
from datetime import datetime
from pathlib import Path
from typing import Any

DOC_OUTPUT_ENV = "TIANSHU_DOC_OUTPUT_DIR"


def doc_output_root() -> Path:
    """Resolve (and create) the directory that ``/api/doc/*`` serves from."""
    raw = os.environ.get(DOC_OUTPUT_ENV)
    root = Path(raw) if raw else Path("/doc_output")
    root.mkdir(parents=True, exist_ok=True)
    return root


def _project_dir(project_id: str | None) -> Path:
    root = doc_output_root()
    if project_id:
        target = root / project_id
        target.mkdir(parents=True, exist_ok=True)
        return target
    return root


def _safe(name: str, limit: int = 40) -> str:
    cleaned = re.sub(r"[^\w一-鿿-]", "_", name or "").strip("_")
    return cleaned[:limit] or "doc"


def _extract_scenario(runtime: Any) -> dict[str, Any]:
    """Pull the current scenario dict out of the runtime export."""
    try:
        exported = runtime.get_exported_scenario()
    except Exception:  # noqa: BLE001 - never let doc gen crash the skill
        return {}
    if isinstance(exported, dict):
        current = exported.get("currentScenario") or exported.get(
            "currentScenarioSnapshot"
        )
        if isinstance(current, dict):
            return current
    return {}


def _unit_lines(units: list[dict[str, Any]], label: str) -> list[str]:
    lines: list[str] = []
    if not units:
        lines.append(f"- 无{label}")
        return lines
    lines.append(f"- {label}共 {len(units)} 个：")
    for unit in units:
        if not isinstance(unit, dict):
            continue
        uid = unit.get("id") or unit.get("name") or "?"
        name = unit.get("name") or unit.get("class_name") or ""
        side = unit.get("side") or unit.get("sideId") or ""
        lat = unit.get("latitude")
        lon = unit.get("longitude")
        coord = (
            f"（{lat:.4f}, {lon:.4f}）"
            if isinstance(lat, (int, float)) and isinstance(lon, (int, float))
            else ""
        )
        line = f"  - `{uid}` {name}".rstrip()
        if side:
            line += f" ［{side}］"
        line += coord
        lines.append(line)
    return lines


def _build_sections(scn: dict[str, Any], title: str, report_type: str) -> list[dict[str, Any]]:
    name = scn.get("name") or "未命名想定"
    current_time = scn.get("currentTime") or scn.get("current_time") or "—"
    sides = scn.get("sides") or []
    aircraft = scn.get("aircraft") or []
    ships = scn.get("ships") or []
    facilities = scn.get("facilities") or []
    airbases = scn.get("airbases") or []
    missions = scn.get("missions") or []
    obstacles = scn.get("obstacles") or []
    reference_points = scn.get("referencePoints") or scn.get("reference_points") or []

    total_units = (
        len(aircraft) + len(ships) + len(facilities) + len(airbases)
    )

    overview = (
        f"想定名称：{name}\n"
        f"当前推演时间：{current_time}\n"
        f"阵营数量：{len(sides)}　单位总数：{total_units}\n"
        f"空中单位：{len(aircraft)}　水面舰艇：{len(ships)}　"
        f"地面设施：{len(facilities)}　机场/基地：{len(airbases)}\n"
        f"任务数量：{len(missions)}　障碍/约束区：{len(obstacles)}　参考点：{len(reference_points)}"
    )

    side_lines = ["- 无阵营" ] if not sides else [
        f"  - `{s.get('id','?')}` {s.get('name','')}".rstrip() for s in sides
    ]

    sections: list[dict[str, Any]] = [
        {"text": "概览", "level": 1, "body": overview},
        {"text": "阵营", "level": 1, "body": "\n".join(side_lines)},
        {"text": "空中单位", "level": 1, "body": "\n".join(_unit_lines(aircraft, "飞机"))},
        {"text": "水面舰艇", "level": 1, "body": "\n".join(_unit_lines(ships, "舰艇"))},
        {"text": "地面设施", "level": 1, "body": "\n".join(_unit_lines(facilities, "设施"))},
        {"text": "机场与基地", "level": 1, "body": "\n".join(_unit_lines(airbases, "基地"))},
        {"text": "任务", "level": 1, "body": _mission_body(missions)},
        {"text": "障碍与约束区", "level": 1, "body": _obstacle_body(obstacles)},
        {"text": "参考点", "level": 1, "body": _reference_body(reference_points)},
    ]
    return sections


def _mission_body(missions: list[dict[str, Any]]) -> str:
    if not missions:
        return "- 无任务"
    lines = [f"- 共 {len(missions)} 个任务："]
    for m in missions:
        if not isinstance(m, dict):
            continue
        mid = m.get("id") or m.get("name") or "?"
        mname = m.get("name") or ""
        mtype = m.get("type") or ""
        line = f"  - `{mid}` {mname}".rstrip()
        if mtype:
            line += f"（{mtype}）"
        lines.append(line)
    return "\n".join(lines)


def _obstacle_body(obstacles: list[dict[str, Any]]) -> str:
    if not obstacles:
        return "- 无障碍/约束区"
    lines = [f"- 共 {len(obstacles)} 个："]
    for o in obstacles:
        if not isinstance(o, dict):
            continue
        oid = o.get("id") or o.get("name") or "?"
        oname = o.get("name") or ""
        otype = o.get("obstacle_type") or o.get("type") or ""
        line = f"  - `{oid}` {oname}".rstrip()
        if otype:
            line += f"（{otype}）"
        lines.append(line)
    return "\n".join(lines)


def _reference_body(points: list[dict[str, Any]]) -> str:
    if not points:
        return "- 无参考点"
    lines = [f"- 共 {len(points)} 个："]
    for p in points:
        if not isinstance(p, dict):
            continue
        pid = p.get("id") or p.get("name") or "?"
        pname = p.get("name") or ""
        lat = p.get("latitude")
        lon = p.get("longitude")
        coord = (
            f"（{lat:.4f}, {lon:.4f}）"
            if isinstance(lat, (int, float)) and isinstance(lon, (int, float))
            else ""
        )
        line = f"  - `{pid}` {pname}".rstrip() + coord
        lines.append(line)
    return "\n".join(lines)


def _sections_to_markdown(title: str, sections: list[dict[str, Any]]) -> str:
    parts = [f"# {title}", ""]
    for s in sections:
        level = int(s.get("level", 1))
        parts.append(f"{'#' * level} {s.get('text', '')}")
        body = s.get("body", "")
        if body:
            parts.append(body)
            parts.append("")
    parts.append("---")
    parts.append(f"*由天枢平台 doc-tools 于 {datetime.now().strftime('%Y-%m-%d %H:%M:%S')} 自动生成*")
    return "\n".join(parts)


def _make_docx(title: str, sections: list[dict[str, Any]], output_path: Path) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    heading = doc.add_heading(title, level=0)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()
    for s in sections:
        doc.add_heading(str(s.get("text", "")), level=min(int(s.get("level", 1)), 4))
        body = s.get("body", "")
        for line in str(body).split("\n"):
            if line.strip():
                doc.add_paragraph(line.strip())
        doc.add_paragraph()
    doc.save(str(output_path))


def _file_entry(path: Path, fmt: str, project_id: str) -> dict[str, Any]:
    url = f"/api/doc/download/{path.name}"
    if project_id:
        url += f"?scenario_id={project_id}"
    return {
        "filename": path.name,
        "format": fmt,
        "path": str(path),
        "size_kb": round(path.stat().st_size / 1024, 1),
        "download_url": url,
    }


def generate_document(
    runtime: Any,
    *,
    report_type: str = "situation",
    title: str = "",
    format: str = "docx",
    project_id: str = "",
    scenario_id: str = "",
    **kwargs: Any,
) -> dict[str, Any]:
    """Generate a situation report document from the live runtime scenario.

    Returns a dict carrying the generated file(s), a human-readable Markdown
    body (so the chat layer can surface it) and a download URL that the
    existing ``/api/doc/download`` route serves.
    """
    scn = _extract_scenario(runtime)
    name = scn.get("name") or "未命名想定"
    title = title or f"{name} 态势报告"
    sections = _build_sections(scn, title, report_type)

    pid = project_id or scenario_id or ""
    pdir = _project_dir(pid or None)
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    base = f"{_safe(title)}_{ts}"

    md_path = pdir / f"{base}.md"
    md_text = _sections_to_markdown(title, sections)
    md_path.write_text(md_text, encoding="utf-8")

    result: dict[str, Any] = {
        "ok": True,
        "title": title,
        "report_type": report_type,
        "project_id": pid,
        "scenario_name": name,
        "sections": len(sections),
        "markdown": md_text,
        "files": [
            _file_entry(md_path, "markdown", pid),
        ],
    }

    primary: dict[str, Any] | None = None
    if format in ("docx", "both"):
        docx_path = pdir / f"{base}.docx"
        try:
            _make_docx(title, sections, docx_path)
            entry = _file_entry(docx_path, "docx", pid)
            result["files"].append(entry)
            primary = entry
        except Exception as exc:  # noqa: BLE001 - degrade gracefully to markdown
            result.setdefault("warnings", []).append(f"docx 生成失败，仅输出 Markdown：{exc}")

    if primary is None:
        primary = result["files"][0]

    result["file"] = primary["filename"]
    result["format"] = primary["format"]
    result["path"] = primary["path"]
    result["download_url"] = primary["download_url"]
    result["size_kb"] = primary["size_kb"]
    result["summary"] = (
        f"已生成态势报告：{title}（{result['format']}，{result['sections']} 节），"
        "文件已保存到当前项目的 AI_Output 文档中心。"
    )
    return result
